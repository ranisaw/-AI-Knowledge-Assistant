"""
Export Services for Word (.docx) and Excel (.xlsx).
Enables onboarding engineers to export answers and verifiable citations.
"""
import io
import time
from typing import List, Dict, Any, Optional
from pathlib import Path


def export_to_docx(question: str, answer: str, citations: List[Dict[str, Any]]) -> bytes:
    """
    Generate a formatted Microsoft Word (.docx) report with question, answer, and citations table.
    """
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
    except ImportError:
        # Fallback text representation if library is missing
        text_content = f"AI KNOWLEDGE ASSISTANT REPORT\n\nQuestion:\n{question}\n\nAnswer:\n{answer}\n\nCitations:\n"
        for c in citations:
            text_content += f"- [Doc {c.get('doc_number')}] {c.get('title')} ({c.get('source')}): {c.get('excerpt')}\n"
        return text_content.encode("utf-8")

    doc = Document()

    # Header title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("AI Knowledge Assistant — Verification Brief")
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42)

    # Sub-header / metadata
    meta_p = doc.add_paragraph()
    meta_p.add_run(f"Generated: {time.strftime('%Y-%m-%d %H:%M:%S UTC', time.gmtime())} | Grounded RAG Pipeline\n").italic = True

    # Question section
    q_heading = doc.add_heading("Inquiry", level=2)
    q_p = doc.add_paragraph(question)
    q_p.style = "Intense Quote"

    # Answer section
    doc.add_heading("Grounded Response", level=2)
    ans_p = doc.add_paragraph(answer)

    # Citations section
    doc.add_heading("Document Citations & Evidence", level=2)
    if not citations:
        doc.add_paragraph("No specific internal documents were cited for this response.")
    else:
        table = doc.add_table(rows=1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Light Shading Accent 1" if "Light Shading Accent 1" in [s.name for s in doc.styles] else "Table Grid"

        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Ref #"
        hdr_cells[1].text = "Document Title"
        hdr_cells[2].text = "Source File"
        hdr_cells[3].text = "Evidence Excerpt"

        for c in citations:
            row_cells = table.add_row().cells
            row_cells[0].text = f"Doc {c.get('doc_number', '-')}"
            row_cells[1].text = str(c.get('title', ''))
            row_cells[2].text = str(c.get('source', ''))
            row_cells[3].text = str(c.get('excerpt', ''))[:200] + "..."

    # Policy footer
    doc.add_paragraph("\n")
    footer_p = doc.add_paragraph()
    footer_run = footer_p.add_run("Note: Generated for internal engineering onboarding. Verify critical operational procedures with primary system leads.")
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor(100, 116, 139)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_to_xlsx(question: str, answer: str, citations: List[Dict[str, Any]]) -> bytes:
    """
    Generate a formatted Microsoft Excel (.xlsx) workbook with Q&A summary and citation breakdown.
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        # Fallback CSV representation
        csv_content = f"Type,Value\nQuestion,\"{question}\"\nAnswer,\"{answer}\"\n"
        for c in citations:
            csv_content += f"Citation,\"{c.get('title')} ({c.get('source')})\"\n"
        return csv_content.encode("utf-8")

    wb = openpyxl.Workbook()

    # Sheet 1: Summary
    ws_summary = wb.active
    ws_summary.title = "Q&A Summary"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")

    ws_summary.append(["Field", "Details"])
    ws_summary.append(["Timestamp", time.strftime("%Y-%m-%d %H:%M:%S UTC", time.gmtime())])
    ws_summary.append(["Engineer Question", question])
    ws_summary.append(["Synthesized Answer", answer])
    ws_summary.append(["Total Citations Cited", len(citations)])

    for col in ["A", "B"]:
        cell = ws_summary[f"{col}1"]
        cell.font = header_font
        cell.fill = header_fill

    ws_summary.column_dimensions["A"].width = 25
    ws_summary.column_dimensions["B"].width = 90

    # Sheet 2: Citations
    ws_cites = wb.create_sheet(title="Citations Breakdown")
    headers = ["Ref #", "Document Title", "Source File", "Doc Type", "Similarity Score", "Chunk ID", "Excerpt"]
    ws_cites.append(headers)

    for col_idx, h in enumerate(headers, 1):
        cell = ws_cites.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill

    for c in citations:
        ws_cites.append([
            f"Doc {c.get('doc_number')}",
            c.get("title", ""),
            c.get("source", ""),
            c.get("doc_type", ""),
            c.get("score", 0.0),
            c.get("chunk_id", ""),
            c.get("excerpt", "")
        ])

    col_widths = [10, 30, 25, 15, 18, 18, 70]
    for idx, width in enumerate(col_widths, 1):
        col_letter = openpyxl.utils.get_column_letter(idx)
        ws_cites.column_dimensions[col_letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
